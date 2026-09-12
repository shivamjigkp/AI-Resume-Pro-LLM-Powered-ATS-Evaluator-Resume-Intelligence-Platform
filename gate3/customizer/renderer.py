"""Resume rendering - generates PDF and DOCX from structured data.

Uses Jinja2 + WeasyPrint for PDF and python-docx for DOCX output.
Both formats are generated for each application:
  - PDF for direct emails and human readability
  - DOCX for ATS compatibility (especially Workday/Taleo)
"""

from pathlib import Path

import jinja2
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from ..config import settings
from ..models import ResumeData


class ResumeRenderer:
    """Renders ResumeData to PDF and DOCX formats."""

    def __init__(self):
        self.template_dir = settings.templates_dir
        self.jinja_env = jinja2.Environment(
            loader=jinja2.FileSystemLoader(str(self.template_dir)),
            autoescape=True,
        )

    def render_pdf(self, resume: ResumeData, output_path: Path) -> Path | None:
        """Render resume to PDF using HTML template + WeasyPrint."""
        try:
            from weasyprint import HTML

            template = self.jinja_env.get_template("resume.html")
            html_content = template.render(resume=resume)

            output_path = Path(output_path)
            output_path.parent.mkdir(parents=True, exist_ok=True)

            HTML(string=html_content).write_pdf(str(output_path))
            return output_path
        except Exception as e:
            return None

    def render_docx(self, resume: ResumeData, output_path: Path) -> Path:
        """Render resume to DOCX for ATS compatibility.

        Uses single-column layout, standard fonts, and clean formatting
        to maximize ATS parsing accuracy.
        """
        doc = Document()

        # Set default font
        style = doc.styles["Normal"]
        font = style.font
        font.name = "Calibri"
        font.size = Pt(11)

        # Reduce margins for more space
        for section in doc.sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.7)
            section.right_margin = Inches(0.7)

        # Name
        name_para = doc.add_paragraph()
        name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        name_run = name_para.add_run(resume.contact.name)
        name_run.font.size = Pt(18)
        name_run.bold = True

        # Contact info
        contact_parts = []
        if resume.contact.email:
            contact_parts.append(resume.contact.email)
        if resume.contact.phone:
            contact_parts.append(resume.contact.phone)
        if resume.contact.location:
            contact_parts.append(resume.contact.location)
        if resume.contact.linkedin:
            contact_parts.append(resume.contact.linkedin)
        if resume.contact.github:
            contact_parts.append(resume.contact.github)

        if contact_parts:
            contact_para = doc.add_paragraph()
            contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            contact_run = contact_para.add_run(" | ".join(contact_parts))
            contact_run.font.size = Pt(9)

        # Summary
        if resume.summary:
            doc.add_heading("Professional Summary", level=2)
            doc.add_paragraph(resume.summary)

        # Experience
        if resume.experience:
            doc.add_heading("Professional Experience", level=2)
            for exp in resume.experience:
                title_para = doc.add_paragraph()
                title_run = title_para.add_run(f"{exp.title} — {exp.company}")
                title_run.bold = True
                date_run = title_para.add_run(
                    f"\n{exp.location}  |  {exp.start_date} – {exp.end_date}"
                )
                date_run.font.size = Pt(9)

                for bullet in exp.bullets:
                    bullet_para = doc.add_paragraph(bullet.text, style="List Bullet")
                    bullet_para.paragraph_format.space_after = Pt(2)

        # Education
        if resume.education:
            doc.add_heading("Education", level=2)
            for edu in resume.education:
                edu_para = doc.add_paragraph()
                edu_run = edu_para.add_run(f"{edu.degree}")
                if edu.field_of_study:
                    edu_run = edu_para.add_run(f" in {edu.field_of_study}")
                edu_run.bold = True
                edu_para.add_run(f" — {edu.institution}")
                if edu.end_date:
                    edu_para.add_run(f"  |  {edu.end_date}")
                if edu.gpa:
                    edu_para.add_run(f"  |  GPA: {edu.gpa}")

        # Skills
        if resume.skills:
            doc.add_heading("Skills", level=2)
            skill_sections = [
                ("Languages", resume.skills.languages),
                ("Frameworks", resume.skills.frameworks),
                ("Tools", resume.skills.tools),
                ("Platforms", resume.skills.platforms),
                ("Certifications", resume.skills.certifications),
            ]
            for label, items in skill_sections:
                if items:
                    para = doc.add_paragraph()
                    label_run = para.add_run(f"{label}: ")
                    label_run.bold = True
                    para.add_run(", ".join(items))

        # Projects
        if resume.projects:
            doc.add_heading("Projects", level=2)
            for project in resume.projects:
                proj_para = doc.add_paragraph()
                proj_run = proj_para.add_run(project.name)
                proj_run.bold = True
                if project.technologies:
                    proj_para.add_run(f" ({', '.join(project.technologies)})")
                doc.add_paragraph(project.description)

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        return output_path

    def render_text(self, resume: ResumeData) -> str:
        """Render resume as plain text (for pasting into forms)."""
        lines = []
        lines.append(resume.contact.name)
        contact_parts = [
            p for p in [
                resume.contact.email,
                resume.contact.phone,
                resume.contact.location,
            ] if p
        ]
        lines.append(" | ".join(contact_parts))
        lines.append("")

        if resume.summary:
            lines.append("PROFESSIONAL SUMMARY")
            lines.append(resume.summary)
            lines.append("")

        if resume.experience:
            lines.append("PROFESSIONAL EXPERIENCE")
            for exp in resume.experience:
                lines.append(f"{exp.title} — {exp.company}")
                lines.append(f"{exp.location} | {exp.start_date} – {exp.end_date}")
                for bullet in exp.bullets:
                    lines.append(f"  • {bullet.text}")
                lines.append("")

        if resume.education:
            lines.append("EDUCATION")
            for edu in resume.education:
                line = f"{edu.degree}"
                if edu.field_of_study:
                    line += f" in {edu.field_of_study}"
                line += f" — {edu.institution}"
                if edu.end_date:
                    line += f" ({edu.end_date})"
                lines.append(line)
            lines.append("")

        if resume.skills:
            lines.append("SKILLS")
            for label, items in [
                ("Languages", resume.skills.languages),
                ("Frameworks", resume.skills.frameworks),
                ("Tools", resume.skills.tools),
            ]:
                if items:
                    lines.append(f"{label}: {', '.join(items)}")

        return "\n".join(lines)
